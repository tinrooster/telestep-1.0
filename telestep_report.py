"""
telestep_report.py - turn a Session into Markdown / DOCX / HTML / PDF / plain text.

Annotated images are re-rendered from the unannotated raw screenshots every time, so the
numbers on the pictures always match the step numbers after you delete, reorder or exclude steps.
"""
import base64
import contextlib
import datetime as dt
import html
import math
import re
import shutil
import traceback
from dataclasses import asdict, dataclass, fields
from pathlib import Path

from PIL import Image

from telestep_core import annotate, clean_text, draw_marks, focus_blur
from telestep_model import environment_text

CALLOUT = {  # level -> (label, fill colour, accent colour)
    "note": ("Note", "E3F2FD", "1E88E5"),
    "warning": ("Warning", "FFF3E0", "F59E0B"),
    "tip": ("Tip", "E8F5E9", "2E9E4F"),
}
STYLES = {"report": "Technical report (all details)",
          "docs": "Documentation (clean, for end users)",
          "onesheet": "Documentation - one sheet (condensed)"}
FORMATS = {"md": "Markdown", "docx": "Word (DOCX)", "html": "HTML (single file)", "pdf": "PDF", "txt": "Plain text",
           "runbook": "Runbook checklist", "zip": "ZIP bundle", "player": "Narrated player (HTML)"}
OUTPUT_FILES = {"md": "steps.md", "docx": "steps.docx", "html": "steps.html", "pdf": "steps.pdf", "txt": "steps.txt",
                "runbook": "runbook.md", "zip": "steps_bundle.zip", "player": "walkthrough.html"}
_GENERATED = re.compile(r"^step_\d{3}_[0-9a-f]{8}(_after)?\.jpg$")


@dataclass
class Template:
    wording: str = "friendly"  # friendly | technical
    header: str = ""  # shown at the top of every page / the page header
    footer: str = ""
    logo: str = ""  # path to a PNG/JPG
    author: str = ""
    accent: str = "#1e88e5"
    show_time: bool = True
    show_window: bool = True
    show_app: bool = True
    environment: bool = True  # appendix with machine/OS/display details
    image_width_in: float = 6.0
    style: str = "report"  # report | docs | onesheet  (see STYLES)
    focus: str = "off"  # off | window | control: blur everything but the clicked window / control
    verified: str = ""  # runbook: "last verified" date (blank = the day it is exported)

    @property
    def docs(self):  # end-user documentation: no timestamps, window/app details, dialog dumps or environment
        return self.style in ("docs", "onesheet")

    @property
    def onesheet(self):
        return self.style == "onesheet"

    @property
    def show_env(self):
        return bool(self.environment) and not self.docs

    def to_dict(self):
        return asdict(self)

    @classmethod
    def from_settings(cls, settings):
        d = dict((settings or {}).get("template", {}))
        if "wording" in (settings or {}):
            d.setdefault("wording", settings["wording"])
        names = {f.name for f in fields(cls)}
        t = cls(**{k: v for k, v in d.items() if k in names})
        if t.style not in STYLES:
            t.style = "report"
        if t.focus not in ("off", "window", "control"):
            t.focus = "off"
        return t


# ---------------------------------------------------------------------------
# images + model
# ---------------------------------------------------------------------------
def render_step_image(session, step, number, focus="off"):
    """Annotated PIL image for a step (None if it has no picture).  Used by export and the review editor.
    focus = 'window' / 'control' blurs everything else (the saved raw screenshot is never changed)."""
    if step.raw and (session.dir / step.raw).exists():
        img = Image.open(session.dir / step.raw).convert("RGB")
        img = focus_blur(img, step.marker, focus)
        img = annotate(img, step.marker, number if step.numbered else None)
        return draw_marks(img, step.marks)
    if step.image and (session.dir / step.image).exists():  # imported v1 session: already annotated
        return Image.open(session.dir / step.image).convert("RGB")
    return None


def _is_auto(step):
    """Steps Telestep made by itself (admin-window warning, 'a dialog appeared'), not the person's own steps/notes."""
    if step.kind == "dialog" or getattr(step, "auto", False):
        return True
    return step.kind == "note" and step.level == "warning" and "runs as administrator" in (step.friendly or "")


@contextlib.contextmanager
def _style_filter(session, tpl):
    """Documentation styles leave out the automatic technical entries.  They are flagged 'skipped' for the duration so
    step numbers (and the numbers drawn on the pictures) stay consecutive; the flags are always restored."""
    flipped = []
    if tpl is not None and tpl.docs:
        for st in session.steps:
            if not st.skip and _is_auto(st):
                st.skip = True
                flipped.append(st)
    try:
        yield
    finally:
        for st in flipped:
            st.skip = False


def render_images(session, tpl=None):
    """(Re)write images/step_NNN_<id>.jpg for every exported step and drop stale generated files."""
    session.make_dirs()
    with _style_filter(session, tpl):
        keep = _render_images(session, tpl.focus if tpl else "off")
    for f in session.img_dir.iterdir():
        if _GENERATED.match(f.name) and f.name not in keep:
            try:
                f.unlink()
            except OSError:
                pass


def _render_images(session, focus="off"):
    nums = session.numbering()
    keep = set()
    for s in session.steps:
        if s.skip or s.kind in ("note", "section"):
            continue
        n = nums.get(s.id, 0)
        img = render_step_image(session, s, n, focus)
        if img is not None and s.raw:
            rel = f"images/step_{n:03d}_{s.id}.jpg"
            img.save(session.dir / rel, "JPEG", quality=88)
            s.image = rel
            keep.add(Path(rel).name)
        if s.raw_after and (session.dir / s.raw_after).exists():
            rel = f"images/step_{n:03d}_{s.id}_after.jpg"
            Image.open(session.dir / s.raw_after).convert("RGB").save(session.dir / rel, "JPEG", quality=88)
            s.image_after = rel
            keep.add(Path(rel).name)
    return keep


def build_blocks(session, tpl):
    out = []
    with _style_filter(session, tpl):
        nums = session.numbering()
        for s in session.steps:
            if s.skip:
                continue
            cap = clean_text(s.caption(tpl.wording))
            if s.kind == "section":
                out.append({"type": "section", "text": cap})
            elif s.kind == "note":
                out.append({"type": "callout", "level": s.level if s.level in CALLOUT else "note", "text": cap})
            else:
                meta = []
                if not tpl.docs:
                    if tpl.show_window and s.window:
                        meta.append(f"Window: {clean_text(s.window)}")
                    if tpl.show_app and s.app:
                        meta.append("App: " + clean_text(s.app + (f" {s.app_version}" if s.app_version else "")))
                    if tpl.show_time and s.when:
                        meta.append(f"Time: {s.when}")
                imgs = []
                if s.image:
                    imgs.append(("Before" if s.image_after else "", s.image))
                if s.image_after:
                    imgs.append(("After", s.image_after))
                out.append({"type": "step", "n": nums.get(s.id, 0), "text": cap, "meta": " | ".join(meta),
                            "images": imgs, "dialog": "" if tpl.docs else clean_text(s.dialog_text),
                            "comment": clean_text(getattr(s, "comment", "") or ""),
                            "audio": s.audio if getattr(s, "audio", "") and (session.dir / s.audio).exists() else "",
                            "path": getattr(s, "path", None) or []})
    return out


def _created(session):
    try:
        return dt.datetime.fromisoformat(str(session.created)[:19])
    except (ValueError, TypeError):
        return None


def _subtitle(session, tpl):
    d = _created(session)
    if tpl.docs:  # no machine name or clock time in user-facing documents
        return f"{d:%B} {d.day}, {d.year}" if d else ""
    return f"Recorded {str(session.created).replace('T', ' ')[:16]} on {clean_text(session.host)}"


def _byline(tpl):
    return f"Prepared by {clean_text(tpl.author)}" if tpl.author else ""


def _logo(tpl):
    return tpl.logo if tpl.logo and Path(tpl.logo).exists() else None


def _cols(n_steps):
    return 3 if n_steps == 3 else 2 if n_steps <= 6 else 3 if n_steps <= 12 else 4


def _md_escape_lines(text):
    return text.strip().splitlines() or [""]


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------
def write_markdown(session, tpl):
    blocks = build_blocks(session, tpl)
    has_sections = any(b["type"] == "section" for b in blocks)
    lines = []
    logo = _logo(tpl)
    if logo:
        try:
            session.make_dirs()
            rel = "images/logo" + (Path(logo).suffix.lower() or ".png")
            shutil.copyfile(logo, session.dir / rel)
            lines += [f"![]({rel})", ""]
        except OSError:
            pass
    if tpl.header:
        lines += [f"_{clean_text(tpl.header)}_", ""]
    lines += [f"# {clean_text(session.title)}", ""]
    if _byline(tpl):
        lines += [f"**{_byline(tpl)}**", ""]
    if _subtitle(session, tpl):
        lines += [f"_{_subtitle(session, tpl)}_", ""]
    for b in blocks:
        if b["type"] == "section":
            lines += [f"## {b['text']}", ""]
        elif b["type"] == "callout":
            label = CALLOUT[b["level"]][0]
            body = _md_escape_lines(b["text"])
            lines.append(f"> **{label}:** {body[0]}")
            lines += [f"> {ln}" for ln in body[1:]]
            lines.append("")
        elif tpl.docs:
            lines += [f"**{b['n']}. {b['text'].replace('`', '')}**", ""]
            if b["comment"]:
                lines += [f"_{b['comment']}_", ""]
            if b["audio"]:
                lines += [f"[Voice note]({b['audio']})", ""]
            for label, rel in b["images"]:
                lines += [f"![Step {b['n']}]({rel})", ""]
        else:
            lines.append(f"{'###' if has_sections else '##'} Step {b['n']}")
            lines.append(f"**{b['text'].replace('`', '')}**  ")
            if b["comment"]:
                lines.append(f"{b['comment']}  ")
            if b["audio"]:
                lines.append(f"[Voice note]({b['audio']})  ")
            if b["meta"]:
                lines.append(f"_{b['meta']}_")
            lines.append("")
            if b["dialog"]:
                lines += ["> **Dialog text:**"] + [f"> {ln}" for ln in _md_escape_lines(b["dialog"])] + [""]
            for label, rel in b["images"]:
                if label:
                    lines += [f"_{label}_", ""]
                lines += [f"![Step {b['n']}]({rel})", ""]
    if tpl.show_env and session.environment:
        lines += ["## Environment", "", "| | |", "|---|---|"]
        lines += [f"| {k} | {str(v).replace('|', '/')} |" for k, v in session.environment.items()]
        lines.append("")
    if tpl.footer:
        lines += ["---", f"_{clean_text(tpl.footer)}_", ""]
    path = session.dir / "steps.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


# ---------------------------------------------------------------------------
# plain text (also used for "copy steps to clipboard")
# ---------------------------------------------------------------------------
def steps_text(session, tpl):
    lines = []
    if tpl.header:
        lines += [clean_text(tpl.header), ""]
    lines += [clean_text(session.title), "=" * min(60, max(8, len(session.title)))]
    if _byline(tpl):
        lines.append(_byline(tpl))
    if _subtitle(session, tpl):
        lines.append(_subtitle(session, tpl))
    lines.append("")
    for b in build_blocks(session, tpl):
        if b["type"] == "section":
            lines += ["", b["text"].upper(), "-" * len(b["text"])]
        elif b["type"] == "callout":
            body = _md_escape_lines(b["text"])
            lines.append(f"   {CALLOUT[b['level']][0]}: {body[0]}")
            lines += [f"   {ln}" for ln in body[1:]]
        else:
            lines.append(f"{b['n']}. {b['text']}")
            if b["comment"]:
                lines.append(f"   {b['comment']}")
            if b["dialog"]:
                lines += [f"   Dialog text: {ln}" for ln in _md_escape_lines(b["dialog"])]
    if tpl.footer:
        lines += ["", "-" * 20, clean_text(tpl.footer)]
    return "\n".join(lines) + "\n"


def write_text(session, tpl):
    path = session.dir / "steps.txt"
    path.write_text(steps_text(session, tpl), encoding="utf-8")
    return path


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def _shade(cell, fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _no_split(row):
    from docx.oxml import OxmlElement
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _page_number(paragraph, size=None):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = paragraph.add_run()
    if size:
        run.font.size = size
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _docx_callout(doc, b, RGBColor):
    label, fill, accent = CALLOUT[b["level"]]
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    _shade(cell, fill)
    p = cell.paragraphs[0]
    lead = p.add_run(f"{label}: ")
    lead.bold = True
    lead.font.color.rgb = RGBColor.from_string(accent)
    p.add_run(b["text"].strip())  # "\n" becomes a line break in Word
    doc.add_paragraph()


def _docx_grid(doc, session, steps, cols, Inches, Pt, max_h_in=9.0):
    """One-sheet layout: steps side by side in a borderless table (each cell = caption + picture)."""
    colw = 7.1 / cols
    t = doc.add_table(rows=0, cols=cols)
    t.autofit = False
    for i in range(0, len(steps), cols):
        cap_row, pic_row = t.add_row(), t.add_row()  # captions on one row, pictures on the next: same baseline
        _no_split(cap_row)
        _no_split(pic_row)
        for c, pc, b in zip(cap_row.cells, pic_row.cells, steps[i:i + cols]):
            c.width = pc.width = Inches(colw)
            p = c.paragraphs[0]
            p.paragraph_format.keep_with_next = True
            r = p.add_run(f"{b['n']}. {b['text'].replace('`', '')}")
            r.bold = True
            r.font.size = Pt(9 if cols <= 3 else 8)
            for _label, rel in b["images"][:1]:
                with Image.open(session.dir / rel) as im:
                    aspect = im.width / im.height
                pc.paragraphs[0].add_run().add_picture(str(session.dir / rel),
                                                       width=Inches(min(colw - 0.2, max_h_in * aspect)))
    doc.add_paragraph()


def write_docx(session, tpl):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    sec = doc.sections[0]
    if tpl.onesheet:
        sec.left_margin = sec.right_margin = Inches(0.6)
        sec.top_margin = sec.bottom_margin = Inches(0.6)
    grey = RGBColor(0x55, 0x55, 0x55)
    # ---- header / footer ----
    hp = sec.header.paragraphs[0]
    if _logo(tpl):
        try:
            hp.add_run().add_picture(_logo(tpl), height=Inches(0.55))
            hp.add_run("   ")
        except Exception:
            pass
    if tpl.header:
        r = hp.add_run(clean_text(tpl.header))
        r.font.size = Pt(10)
        r.font.color.rgb = grey
    fp = sec.footer.paragraphs[0]
    if tpl.footer:
        r = fp.add_run(clean_text(tpl.footer) + "    ")
        r.font.size = Pt(10)
        r.font.color.rgb = grey
    fp.add_run("Page ").font.size = Pt(10)
    _page_number(fp, Pt(10))

    doc.add_heading(clean_text(session.title), 0)
    if _byline(tpl):
        r = doc.add_paragraph().add_run(_byline(tpl))
        r.bold = True
        r.font.size = Pt(11)
    if _subtitle(session, tpl):
        doc.add_paragraph().add_run(_subtitle(session, tpl)).italic = True

    blocks = build_blocks(session, tpl)
    n_steps = sum(1 for b in blocks if b["type"] == "step")
    cols = _cols(n_steps)
    max_h_in = max(6.0 / max(1, math.ceil(n_steps / cols)) - 0.45, 1.2)  # aim for one page when it fits
    pending = []

    def flush():
        if pending:
            _docx_grid(doc, session, pending, cols, Inches, Pt, max_h_in)
            pending.clear()

    for b in blocks:
        if tpl.onesheet and b["type"] == "step":
            pending.append(b)
            continue
        flush()
        if b["type"] == "section":
            doc.add_heading(b["text"], level=2 if tpl.onesheet else 1)
        elif b["type"] == "callout":
            _docx_callout(doc, b, RGBColor)
        elif tpl.docs:
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(10)
            n = p.add_run(f"{b['n']}.  ")
            n.bold = True
            n.font.color.rgb = RGBColor.from_string((tpl.accent or "#1e88e5").lstrip("#"))
            p.add_run(b["text"].replace("`", "")).bold = True
            if b["comment"]:
                doc.add_paragraph().add_run(b["comment"]).italic = True
            for _label, rel in b["images"]:
                doc.add_picture(str(session.dir / rel), width=Inches(tpl.image_width_in))
        else:
            doc.add_heading(f"Step {b['n']}", level=2)
            run = doc.add_paragraph().add_run(b["text"].replace("`", ""))
            run.bold = True
            if b["comment"]:
                doc.add_paragraph().add_run(b["comment"]).italic = True
            if b["meta"]:
                m = doc.add_paragraph().add_run(b["meta"])
                m.italic = True
                m.font.size = Pt(9)
            if b["dialog"]:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.add_run("Dialog text: ").bold = True
                p.add_run(b["dialog"])
            for label, rel in b["images"]:
                if label:
                    doc.add_paragraph().add_run(label).italic = True
                doc.add_picture(str(session.dir / rel), width=Inches(tpl.image_width_in))
    flush()
    if tpl.show_env and session.environment:
        doc.add_heading("Environment", level=1)
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        for k, v in session.environment.items():
            row = t.add_row().cells
            row[0].text, row[1].text = clean_text(k), clean_text(str(v))
    path = session.dir / "steps.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# HTML (single self-contained file)
# ---------------------------------------------------------------------------
def _data_uri(path):
    p = Path(path)
    mime = {".png": "image/png", ".wav": "audio/wav"}.get(p.suffix.lower(), "image/jpeg")
    return f"data:{mime};base64," + base64.b64encode(p.read_bytes()).decode("ascii")


def _h(text):
    return html.escape(text).replace("\n", "<br>")


def write_html(session, tpl):
    acc = tpl.accent or "#1e88e5"
    blocks = build_blocks(session, tpl)
    cols = _cols(sum(1 for b in blocks if b["type"] == "step"))
    css = f"""
    :root {{ --acc: {acc}; }}
    body {{ font: 15px/1.5 "Segoe UI", system-ui, sans-serif; color: #1c1f23; max-width: {1100 if tpl.onesheet else 920}px; margin: 0 auto; padding: 24px 20px 60px; }}
    header.top {{ display: flex; align-items: center; gap: 14px; color: #556; font-size: 14px; border-bottom: 1px solid #dde; padding-bottom: 8px; margin-bottom: 18px; }}
    header.top img {{ height: 48px; }}
    h1 {{ margin: 0 0 4px; }} .by {{ font-weight: 600; }} .sub {{ color: #667; font-style: italic; margin-bottom: 22px; }}
    h2.section {{ margin-top: 34px; padding-bottom: 4px; border-bottom: 2px solid var(--acc); }}
    .step {{ margin: 22px 0; page-break-inside: avoid; break-inside: avoid; }}
    .step h3 {{ margin: 0; color: var(--acc); font-size: 14px; text-transform: uppercase; letter-spacing: .05em; }}
    .step .cap {{ font-weight: 600; font-size: 17px; margin: 2px 0; }}
    .step .cap .num {{ color: var(--acc); margin-right: 6px; }}
    .meta {{ color: #778; font-size: 12px; }}
    .comment {{ color: #445; font-style: italic; margin: 2px 0 4px; }}
    audio {{ display: block; width: 100%; max-width: 360px; margin: 4px 0 6px; }}
    figure {{ margin: 8px 0 0; }} figure img {{ max-width: 100%; border: 1px solid #ccd; border-radius: 4px; }}
    figcaption {{ font-size: 12px; color: #778; margin-bottom: 2px; }}
    .callout {{ border-left: 4px solid; padding: 8px 12px; margin: 16px 0; border-radius: 3px; }}
    .callout b {{ margin-right: 4px; }}
    .dialog {{ background: #f6f6f8; border-left: 3px solid #999; padding: 6px 10px; margin: 6px 0; font-family: Consolas, monospace; font-size: 13px; }}
    table.env {{ border-collapse: collapse; font-size: 13px; }} table.env td {{ border: 1px solid #ccd; padding: 3px 10px; }}
    table.env td:first-child {{ color: #556; }}
    footer {{ margin-top: 40px; color: #556; font-size: 13px; border-top: 1px solid #dde; padding-top: 8px; }}
    .grid {{ display: grid; grid-template-columns: repeat({cols}, 1fr); gap: 14px; align-items: start; }}
    .grid .full {{ grid-column: 1 / -1; }}
    .grid .step {{ margin: 0; display: grid; grid-template-rows: subgrid; grid-row: span 2; row-gap: 4px; }}
    .grid .step .cap {{ font-size: 13px; }} .grid .step figure {{ margin: 0; }}
    @page {{ margin: 12mm; }}
    @media print {{ body {{ max-width: none; padding: 0; }} }}
    """
    out = [f"<!doctype html><html lang='en'><head><meta charset='utf-8'><title>{html.escape(session.title)}</title>"
           f"<meta name='viewport' content='width=device-width,initial-scale=1'><style>{css}</style></head><body>"]
    if tpl.header or _logo(tpl):
        logo = f"<img alt='' src='{_data_uri(_logo(tpl))}'>" if _logo(tpl) else ""
        out.append(f"<header class='top'>{logo}<span>{_h(clean_text(tpl.header))}</span></header>")
    out.append(f"<h1>{_h(clean_text(session.title))}</h1>")
    if _byline(tpl):
        out.append(f"<div class='by'>{_h(_byline(tpl))}</div>")
    out.append(f"<div class='sub'>{_h(_subtitle(session, tpl))}</div>")
    if tpl.onesheet:
        out.append("<div class='grid'>")
    for b in blocks:
        full = " full" if tpl.onesheet else ""
        if b["type"] == "section":
            out.append(f"<h2 class='section{full}'>{_h(b['text'])}</h2>")
        elif b["type"] == "callout":
            label, fill, accent = CALLOUT[b["level"]]
            out.append(f"<div class='callout{full}' style='background:#{fill};border-color:#{accent}'>"
                       f"<b style='color:#{accent}'>{label}:</b>{_h(b['text'].strip())}</div>")
        elif tpl.docs:
            out.append(f"<section class='step'><div class='cap'><span class='num'>{b['n']}.</span>{_h(b['text'])}</div>")
            if b["comment"] and not tpl.onesheet:
                out.append(f"<div class='comment'>{_h(b['comment'])}</div>")
            if b["audio"] and not tpl.onesheet:
                out.append(f"<audio controls preload='none' src='{_data_uri(session.dir / b['audio'])}'></audio>")
            for label, rel in (b["images"][:1] if tpl.onesheet else b["images"]):
                out.append(f"<figure><img alt='Step {b['n']}' src='{_data_uri(session.dir / rel)}'></figure>")
            out.append("</section>")
        else:
            out.append(f"<section class='step'><h3>Step {b['n']}</h3><div class='cap'>{_h(b['text'])}</div>")
            if b["comment"]:
                out.append(f"<div class='comment'>{_h(b['comment'])}</div>")
            if b["audio"]:
                out.append(f"<audio controls preload='none' src='{_data_uri(session.dir / b['audio'])}'></audio>")
            if b["meta"]:
                out.append(f"<div class='meta'>{_h(b['meta'])}</div>")
            if b["dialog"]:
                out.append(f"<div class='dialog'>{_h(b['dialog'])}</div>")
            for label, rel in b["images"]:
                cap = f"<figcaption>{label}</figcaption>" if label else ""
                out.append(f"<figure>{cap}<img alt='Step {b['n']}' src='{_data_uri(session.dir / rel)}'></figure>")
            out.append("</section>")
    if tpl.onesheet:
        out.append("</div>")
    if tpl.show_env and session.environment:
        out.append("<h2 class='section'>Environment</h2><table class='env'>")
        out += [f"<tr><td>{_h(k)}</td><td>{_h(str(v))}</td></tr>" for k, v in session.environment.items()]
        out.append("</table>")
    if tpl.footer:
        out.append(f"<footer>{_h(clean_text(tpl.footer))}</footer>")
    out.append("</body></html>")
    path = session.dir / "steps.html"
    path.write_text("".join(out), encoding="utf-8")
    return path


# ---------------------------------------------------------------------------
# PDF (reportlab, pure Python)
# ---------------------------------------------------------------------------
_FONT_CANDIDATES = [
    (r"C:\Windows\Fonts\segoeui.ttf", r"C:\Windows\Fonts\segoeuib.ttf"),
    (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
]


def _pdf_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for reg, bold in _FONT_CANDIDATES:
        if Path(reg).exists() and Path(bold).exists():
            try:
                pdfmetrics.registerFont(TTFont("SR", reg))
                pdfmetrics.registerFont(TTFont("SR-B", bold))
                pdfmetrics.registerFontFamily("SR", normal="SR", bold="SR-B", italic="SR", boldItalic="SR-B")
                return "SR", "SR-B", True
            except Exception:
                continue
    return "Helvetica", "Helvetica-Bold", False


def write_pdf(session, tpl):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image as RLImage
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    font, bold, unicode_ok = _pdf_fonts()

    def T(s):  # text safe for the chosen font
        s = escape(clean_text(s))
        if not unicode_ok:
            s = s.encode("latin-1", "replace").decode("latin-1")
        return s.replace("\n", "<br/>")

    accent = colors.HexColor(tpl.accent or "#1e88e5")
    base = ParagraphStyle("b", fontName=font, fontSize=10.5, leading=14)
    st_title = ParagraphStyle("t", parent=base, fontName=bold, fontSize=22, leading=26, spaceAfter=4)
    st_by = ParagraphStyle("by", parent=base, fontName=bold, fontSize=10.5, spaceAfter=2)
    st_sub = ParagraphStyle("s", parent=base, textColor=colors.HexColor("#667788"), fontSize=9, spaceAfter=12)
    st_sec = ParagraphStyle("sec", parent=base, fontName=bold, fontSize=15, leading=19, spaceBefore=14, spaceAfter=6)
    st_step = ParagraphStyle("st", parent=base, fontName=bold, fontSize=8.5, textColor=accent, spaceBefore=12)
    st_cap = ParagraphStyle("c", parent=base, fontName=bold, fontSize=12, leading=15)
    st_meta = ParagraphStyle("m", parent=base, fontSize=8, textColor=colors.HexColor("#778899"))
    st_comment = ParagraphStyle("cm", parent=base, fontSize=10, textColor=colors.HexColor("#445566"))
    st_lbl = ParagraphStyle("l", parent=base, fontSize=8, textColor=colors.HexColor("#778899"), spaceBefore=4)

    page_w, page_h = A4
    margin = (12 if tpl.onesheet else 18) * mm
    avail_w, avail_h = page_w - 2 * margin, page_h - 2 * margin - 30

    def picture(rel, max_h=None, max_w=None):
        p = session.dir / rel
        with Image.open(p) as im:
            w, h = im.size
        k = min((max_w or avail_w) / w, (max_h or avail_h * (0.6 if tpl.docs else 0.75)) / h)
        return RLImage(str(p), width=w * k, height=h * k)

    def callout(b):
        label, fill, acc = CALLOUT[b["level"]]
        para = Paragraph(f"<font name='{bold}' color='#{acc}'>{label}:</font> {T(b['text'].strip())}", base)
        tbl = Table([[para]], colWidths=[avail_w])
        tbl.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#" + fill)),
                                 ("LINEBEFORE", (0, 0), (0, -1), 3, colors.HexColor("#" + acc)),
                                 ("LEFTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 5),
                                 ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
        return tbl

    story = [Paragraph(T(session.title), st_title)]
    if _byline(tpl):
        story.append(Paragraph(T(_byline(tpl)), st_by))
    story.append(Paragraph(T(_subtitle(session, tpl)), st_sub))
    blocks = build_blocks(session, tpl)

    if tpl.onesheet:
        n_steps = sum(1 for b in blocks if b["type"] == "step")
        cols = _cols(n_steps)
        col_w = avail_w / cols
        n_rows = max(1, math.ceil(n_steps / cols))
        row_h = (page_h - 2 * margin - 40 * mm) / n_rows
        max_h = max(row_h - 12 * mm, 28 * mm)
        st_small = ParagraphStyle("sm", parent=base, fontName=bold, fontSize=9 if cols <= 3 else 8, leading=11)
        rows, spans, cur, cur_pics = [], [], [], []

        def flush_cur():
            if cur:  # captions on one table row, pictures on the next, so every picture starts at the same height
                rows.append(cur + [""] * (cols - len(cur)))
                rows.append(cur_pics + [""] * (cols - len(cur_pics)))
                cur.clear()
                cur_pics.clear()

        for b in blocks:
            if b["type"] == "step":
                cur.append(Paragraph(f"<font color='#{(tpl.accent or '#1e88e5').lstrip('#')}'>{b['n']}.</font> "
                                     f"{T(b['text'].replace('`', ''))}", st_small))
                pic = [picture(rel, max_h=max_h, max_w=col_w - 10) for _label, rel in b["images"][:1]]
                cur_pics.append(pic[0] if pic else "")
                if len(cur) == cols:
                    flush_cur()
            else:
                flush_cur()
                spans.append(("SPAN", (0, len(rows)), (-1, len(rows))))
                rows.append([Paragraph(T(b["text"]), st_sec) if b["type"] == "section" else callout(b)]
                            + [""] * (cols - 1))
        flush_cur()
        if rows:
            tbl = Table(rows, colWidths=[col_w] * cols)
            tbl.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 3),
                                     ("RIGHTPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (-1, -1), 3),
                                     ("BOTTOMPADDING", (0, 0), (-1, -1), 4)] + spans))
            story.append(tbl)
    else:
        for b in blocks:
            if b["type"] == "section":
                story.append(Paragraph(T(b["text"]), st_sec))
            elif b["type"] == "callout":
                story += [Spacer(1, 6), callout(b), Spacer(1, 6)]
            elif tpl.docs:
                part = [Spacer(1, 8), Paragraph(f"<font color='#{(tpl.accent or '#1e88e5').lstrip('#')}'>{b['n']}.</font> "
                                                f"{T(b['text'].replace('`', ''))}", st_cap)]
                if b["comment"]:
                    part.append(Paragraph(f"<i>{T(b['comment'])}</i>", st_comment))
                for _label, rel in b["images"]:
                    part += [Spacer(1, 3), picture(rel)]
                story.append(KeepTogether(part))
            else:
                part = [Paragraph(f"STEP {b['n']}", st_step), Paragraph(T(b["text"]), st_cap)]
                if b["comment"]:
                    part.append(Paragraph(f"<i>{T(b['comment'])}</i>", st_comment))
                if b["meta"]:
                    part.append(Paragraph(T(b["meta"]), st_meta))
                if b["dialog"]:
                    part.append(Paragraph("<i>Dialog text:</i> " + T(b["dialog"]), base))
                for label, rel in b["images"]:
                    if label:
                        part.append(Paragraph(T(label), st_lbl))
                    part += [Spacer(1, 3), picture(rel)]
                    story.append(KeepTogether(part))
                    part = []
                if part:
                    story.append(KeepTogether(part))
    if tpl.show_env and session.environment:
        story.append(Paragraph("Environment", st_sec))
        rows = [[Paragraph(T(k), st_meta), Paragraph(T(str(v)), base)] for k, v in session.environment.items()]
        tbl = Table(rows, colWidths=[avail_w * 0.3, avail_w * 0.7])
        tbl.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#ccccdd")),
                                 ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(tbl)

    logo = _logo(tpl)

    def plain(s):
        return T(s).replace("<br/>", " ").replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")

    def decorate(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#556677"))
        x = margin
        if logo:
            try:
                with Image.open(logo) as im:
                    lw, lh = im.size
                h = 12 * mm
                canvas.drawImage(logo, x, page_h - margin - 1 * mm, width=h * lw / lh, height=h,
                                 preserveAspectRatio=True, mask="auto")
                x += h * lw / lh + 4 * mm
            except Exception:
                pass
        if tpl.header:
            canvas.setFont(font, 10)
            canvas.drawString(x, page_h - margin + 4 * mm, plain(tpl.header)[:110])
        canvas.setFont(font, 9)
        foot = (plain(tpl.footer) + "    ") if tpl.footer else ""
        canvas.drawString(margin, margin - 12, foot + f"Page {doc.page}")
        canvas.restoreState()

    path = session.dir / "steps.pdf"
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=margin, rightMargin=margin,
                      topMargin=margin + (10 * mm if logo else 4), bottomMargin=margin,
                      title=clean_text(session.title)).build(story, onFirstPage=decorate, onLaterPages=decorate)
    return path


# ---------------------------------------------------------------------------
# runbook checklist (Markdown with tick boxes and a verification log)
# ---------------------------------------------------------------------------
def write_runbook(session, tpl):
    """A procedure people can tick off: `- [ ]` per step, notes/warnings as quotes, and a 'last verified' line + log."""
    blocks = build_blocks(session, tpl)
    today = f"{dt.date.today():%Y-%m-%d}"
    when = (tpl.verified or "").strip() or today
    who = clean_text(tpl.author) or "________"
    lines = []
    if tpl.header:
        lines += [f"_{clean_text(tpl.header)}_", ""]
    lines += [f"# {clean_text(session.title)}", "", f"**Runbook** | Last verified: **{when}** by **{who}**", ""]
    notes = [b for b in blocks if b["type"] == "callout" and b["level"] == "warning"]
    if notes:
        lines += ["## Before you start", ""]
        lines += [f"- {n['text'].strip().splitlines()[0] if n['text'].strip() else ''}" for n in notes]
        lines.append("")
    lines += ["## Procedure", ""]
    for b in blocks:
        if b["type"] == "section":
            lines += [f"### {b['text']}", ""]
        elif b["type"] == "callout":
            body = _md_escape_lines(b["text"])
            lines.append(f"> **{CALLOUT[b['level']][0]}:** {body[0]}")
            lines += [f"> {ln}" for ln in body[1:]]
            lines.append("")
        else:
            lines.append(f"- [ ] **{b['n']}.** {b['text'].replace('`', '')}")
            if b["comment"]:
                lines.append(f"  - {b['comment']}")
            if b["images"]:
                lines.append(f"  - Screenshot: [{b['images'][0][1]}]({b['images'][0][1]})")
    lines += ["", "## Verification log", "", "| Date | By | Result | Notes |", "|---|---|---|---|",
              f"| {when} | {who} | | |", "| | | | |", ""]
    if tpl.footer:
        lines += ["---", f"_{clean_text(tpl.footer)}_", ""]
    path = session.dir / "runbook.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


# ---------------------------------------------------------------------------
# ZIP bundle: steps.md plus only the pictures it uses, ready to drop into a wiki or ticket
# ---------------------------------------------------------------------------
def write_bundle(session, tpl):
    import zipfile
    md = write_markdown(session, tpl)
    blocks = [b for b in build_blocks(session, tpl) if b["type"] == "step"]
    used = {rel for b in blocks for _l, rel in b["images"]} | {b["audio"] for b in blocks if b["audio"]}
    logo = _logo(tpl)
    if logo:
        used.add("images/logo" + (Path(logo).suffix.lower() or ".png"))
    path = session.dir / "steps_bundle.zip"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(md, "steps.md")
        for rel in sorted(used):
            f = session.dir / rel
            if f.exists():
                z.write(f, rel.replace("\\", "/"))
    return path


# ---------------------------------------------------------------------------
# entry point
# ---------------------------------------------------------------------------
def write_player(session, tpl):
    import telestep_player
    return telestep_player.write(session, tpl)


WRITERS = {"md": write_markdown, "docx": write_docx, "html": write_html, "pdf": write_pdf, "txt": write_text,
           "runbook": write_runbook, "zip": write_bundle, "player": write_player}


def export(session, tpl, formats):
    """Write the requested formats.  Returns {format: path-string | 'FAILED: ...'}; one failing format
    never stops the others (details go to save_errors.log)."""
    session.settings["wording"] = tpl.wording
    session.settings["template"] = tpl.to_dict()
    results, trace = {}, []
    try:
        render_images(session, tpl)
    except Exception as e:
        results["images"] = f"FAILED: {type(e).__name__}: {e}"
        trace.append("images\n" + traceback.format_exc())
    session.save()
    if session.environment and not tpl.docs:
        try:
            (session.dir / "environment.txt").write_text(environment_text(session.environment, session.title),
                                                         encoding="utf-8")
        except OSError:
            pass
    for fmt in formats:
        w = WRITERS.get(fmt)
        if not w:
            results[fmt] = "FAILED: unknown format"
            continue
        try:
            results[fmt] = str(w(session, tpl))
        except Exception as e:
            results[fmt] = f"FAILED: {type(e).__name__}: {e}"
            trace.append(f"{fmt}\n{traceback.format_exc()}")
    log = session.dir / "save_errors.log"
    if trace:
        try:
            log.write_text(f"{dt.datetime.now():%Y-%m-%d %H:%M:%S}\n\n" + "\n".join(trace), encoding="utf-8")
        except OSError:
            pass
    return results
